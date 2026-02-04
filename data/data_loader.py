import re
from dataclasses import dataclass
from typing import Any, Dict, List, Optional, Tuple

import openpyxl


def norm(s: Any) -> str:
    """Uppercase + collapse to A-Z0-9 spaces."""
    if s is None:
        return ""
    s = str(s).upper()
    s = re.sub(r"[^A-Z0-9]+", " ", s)
    s = re.sub(r"\s+", " ", s).strip()
    return s


def first_token(s: Any) -> str:
    t = norm(s)
    return t.split()[0] if t else ""


def safe_float(x: Any) -> Optional[float]:
    if x is None:
        return None
    if isinstance(x, (int, float)):
        return float(x)
    s = str(x).strip()
    if not s:
        return None
    s = s.replace(",", "")
    try:
        return float(s)
    except ValueError:
        return None


@dataclass
class SprayRow:
    sheet: str
    catchment: Optional[str]
    drain: str
    distance_m: float
    qty_km: float
    start_m: Optional[float] = None
    end_m: Optional[float] = None
    lat: Optional[float] = None
    lon: Optional[float] = None

    work_order: Optional[str] = None
    po: Optional[str] = None

    # Persistent workflow state
    completed: bool = False
    invoiced_at: Optional[str] = None  # e.g. "2026-01-29"


def load_spray_list(path: str) -> List[SprayRow]:
    """
    Reads ALL sheets in Spray list.xlsx.

    Expected layout:
      Column A = Drain Name
      Column D = Distance (m)

    Catchment header rows contain "Catchment" in column A.
    """
    wb = openpyxl.load_workbook(path, data_only=True)
    out: List[SprayRow] = []

    for ws in wb.worksheets:
        current_catchment: Optional[str] = None

        # Try to find header row containing "Drain Name" in column A
        header_row = None
        for r in range(1, min(ws.max_row, 60) + 1):
            a_norm = norm(ws.cell(r, 1).value)
            d_norm = norm(ws.cell(r, 4).value)
            if a_norm == "DRAIN NAME":
                header_row = r
                break
            if "DRAIN" in a_norm and d_norm in {"TOTAL DIST", "TOTAL DISTANCE"}:
                header_row = r
                break

        start_row = (header_row + 1) if header_row else 2

        last_drain: Optional[str] = None

        for r in range(start_row, ws.max_row + 1):
            a = ws.cell(r, 1).value  # Drain Name or Catchment
            b = ws.cell(r, 2).value  # Start (m)
            c = ws.cell(r, 3).value  # End (m)
            d = ws.cell(r, 4).value  # Distance (m)
            lat_raw = ws.cell(r, 6).value  # Pins -> Lat
            lon_raw = ws.cell(r, 7).value  # Pins -> Long

            a_str = str(a).strip() if a is not None else ""
            a_norm = norm(a_str)

            # Skip blank rows
            if not a_str and (d is None or str(d).strip() == ""):
                continue

            # Catchment header
            if a_str and "CATCHMENT" in a_norm and (d is None or str(d).strip() == ""):
                current_catchment = a_str.strip()
                continue

            # Skip headers/totals
            if a_norm in {"DRAIN NAME", "TOTAL DIST", "TOTAL DISTANCE"}:
                continue
            if a_norm.startswith("TOTAL "):
                continue

            drain = a_str.strip()
            if not drain:
                if last_drain:
                    drain = last_drain
                else:
                    continue
            else:
                last_drain = drain

            start_m = safe_float(b)
            end_m = safe_float(c)
            if start_m is not None and end_m is not None:
                def _fmt_m(x: float) -> str:
                    return str(int(x)) if float(x).is_integer() else str(x)
                drain = f"{drain} ({_fmt_m(start_m)}-{_fmt_m(end_m)})"

            dist = safe_float(d)
            if dist is None and start_m is not None and end_m is not None:
                if end_m >= start_m:
                    dist = end_m - start_m
            if dist is None:
                continue

            lat = safe_float(lat_raw)
            lon = safe_float(lon_raw)

            out.append(
                SprayRow(
                    sheet=ws.title,
                    catchment=current_catchment,
                    drain=drain,
                    distance_m=float(dist),
                    qty_km=round(float(dist) / 1000.0, 2),
                    start_m=start_m,
                    end_m=end_m,
                    lat=lat,
                    lon=lon,
                )
            )

    return out


def _find_header_row(ws, max_scan_rows: int = 80, max_scan_cols: int = 120) -> int:
    """Find a header row containing MI and MAINTITEM TEXT (or LOCATION GROUPING)."""
    for r in range(1, min(ws.max_row, max_scan_rows) + 1):
        row = [norm(ws.cell(r, c).value) for c in range(1, max_scan_cols + 1)]
        if ("MI" in row or "MI #" in row) and (
            "MAINTITEM TEXT" in row or "LOCATION GROUPING" in row
        ):
            return r
    return 2


def _group_after_prefix(maint_text: Any) -> str:
    """
    From MaintItem text like:
      "52W SPRAY DRAINS BENGER ..."
    return "BENGER".
    """
    t = norm(maint_text)
    if not t:
        return ""

    toks = t.split()
    prefix = ["52W", "SPRAY", "DRAINS"]

    for i in range(0, max(0, len(toks) - len(prefix)) + 1):
        if toks[i : i + len(prefix)] == prefix:
            j = i + len(prefix)
            return toks[j] if j < len(toks) else ""

    return ""


def load_spray_drains_mapping(
    path: str, sheet_name: str = "Spray Drains"
) -> Tuple[str, Dict[str, str]]:
    """
    Reads ONLY the 'Spray Drains' sheet (or falls back to first sheet).
    Builds mapping:
      GROUP (e.g. BENGER) -> MI work order
    PO read from row 1 (A1/B1/C1 fallback).
    """
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb[sheet_name] if sheet_name in wb.sheetnames else wb.worksheets[0]

    po_raw = ws["A1"].value or ws["B1"].value or ws["C1"].value
    po = str(po_raw).strip() if po_raw is not None else ""

    header_row = _find_header_row(ws)

    mi_col: Optional[int] = None
    maint_col: Optional[int] = None

    for c in range(1, 120):
        h = norm(ws.cell(header_row, c).value)
        if h in {"MI", "MI #"}:
            mi_col = c
        elif h == "MAINTITEM TEXT":
            maint_col = c

    # fallback to Location grouping if needed
    if maint_col is None:
        for c in range(1, 120):
            if norm(ws.cell(header_row, c).value) == "LOCATION GROUPING":
                maint_col = c
                break

    if mi_col is None:
        mi_col = 2
    if maint_col is None:
        return po, {}

    mapping: Dict[str, str] = {}

    for r in range(header_row + 1, ws.max_row + 1):
        mi = ws.cell(r, mi_col).value
        maint = ws.cell(r, maint_col).value
        if not mi or not maint:
            continue

        mi_s = str(mi).strip()
        if not mi_s:
            continue

        group = _group_after_prefix(maint)
        if not group:
            continue

        mapping[group] = mi_s  # last wins

    return po, mapping


def attach_work_info(
    spray_rows: List[SprayRow], po: str, mapping: Dict[str, str]
) -> int:
    """
    Match by GROUP = first word of spray drain (e.g., 'BENGER' from 'BENGER C').

    If no match: leave WO/PO blank (likely already claimed).
    """
    matched = 0
    for r in spray_rows:
        group = first_token(r.drain)
        if group and group in mapping:
            r.work_order = mapping[group]
            r.po = po
            matched += 1
        else:
            r.work_order = None
            r.po = None
    return matched


# ==========================
# Work List generic loader
# ==========================


@dataclass
class WorkListRow:
    sheet: str
    mp: str
    wo: str
    location: str
    call_date: str
    po: str


def _first_non_empty_in_row(ws, row_num: int) -> str:
    for cell in ws[row_num]:
        v = cell.value
        if v is None:
            continue
        s = str(v).strip()
        if s:
            return s
    return ""


def load_work_list_sheet(workbook_path: str, sheet_name: str) -> List[WorkListRow]:
    """
    Work List sheets:
      - PO somewhere in row 1 (often A1)
      - Header row at 2
      - Data from row 3
      - Column A: MP #
      - Column B: MI # (Work Order)
      - Column C: MaintItem text (common, ignore)
      - Column D: Suburb/Town (location)
      - Column E: Call Date
    """
    wb = openpyxl.load_workbook(workbook_path, data_only=True)
    if sheet_name not in wb.sheetnames:
        return []
    ws = wb[sheet_name]

    po = _first_non_empty_in_row(ws, 1)

    out: List[WorkListRow] = []

    for r in range(3, ws.max_row + 1):
        a = ws.cell(row=r, column=1).value  # MP #
        b = ws.cell(row=r, column=2).value  # MI # (WO)
        d = ws.cell(row=r, column=4).value  # Suburb/Town
        e = ws.cell(row=r, column=5).value  # Call Date

        # Skip blank rows
        if a is None and b is None and d is None and e is None:
            continue

        # Skip repeated header rows inside sheet
        if isinstance(a, str) and a.strip().lower() == "mp #":
            continue

        mp = "" if a is None else str(a).strip()
        wo = "" if b is None else str(b).strip()
        location = "" if d is None else str(d).strip()
        call_date = "" if e is None else str(e).strip()

        if not wo:
            continue

        out.append(
            WorkListRow(
                sheet=sheet_name,
                mp=mp,
                wo=wo,
                location=location,
                call_date=call_date,
                po=po,
            )
        )

    return out
